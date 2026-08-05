# -*- coding: utf-8 -*-
"""误分类报错工具 — 由 first_run.bat 或 CLI 调用
交互式选择邮件 → 输入正确分级 → 写入 state/misclassifications.jsonl
"""
import json
import os
import sys
from datetime import datetime
from pathlib import Path

STATE_DIR = Path('state')
ERROR_FILE = STATE_DIR / 'misclassifications.jsonl'
CYCLE_DIR = Path('reports/cycles')

LEVEL_NAMES = {
    0: '0级-紧急(店铺级)',
    1: '1级-重要(店铺/商品)',
    2: '2级-关注(链接/日常)',
    3: '3级-客服来信',
    4: '4级-日常管理',
    5: '5级-出单/发货',
    6: '6级-无关/广告',
}

SUB_CATS = {
    0: ['店铺关停', 'POA申诉', '重大Case'],
    1: ['系统质检', '商品下架', '认证审核', '账户安全'],
    2: ['系统调整', '资金付款', '货件到仓', '链接合规', '系统提醒', '退款退货', '业务调整'],
    3: ['客服来信'],
    4: ['日常管理'],
    5: ['出单通知', '发货通知'],
    6: ['活动推荐', '广告推销'],
}


def load_recent_emails(max_count=20):
    """从最新的周期报表中读取最近的邮件"""
    if not CYCLE_DIR.exists():
        return []

    # 找最新的 Excel 文件
    excel_files = sorted(CYCLE_DIR.glob('slot_*.xlsx'), key=lambda f: f.stat().st_mtime, reverse=True)
    if not excel_files:
        return []

    try:
        from openpyxl import load_workbook
        wb = load_workbook(excel_files[0], read_only=True)
        ws = wb.active
        emails = []
        for row in range(2, ws.max_row + 1):
            subject = ws.cell(row=row, column=10).value or ''
            sender = ws.cell(row=row, column=9).value or ''
            level = ws.cell(row=row, column=5).value
            sub_cat = ws.cell(row=row, column=6).value or ''
            date = ws.cell(row=row, column=1).value or ''
            owner = ws.cell(row=row, column=4).value or ''
            source = ws.cell(row=row, column=3).value or ''
            try:
                level = int(level) if level is not None else 6
            except (ValueError, TypeError):
                level = 6
            if level > 6:
                level = 6
            emails.append({
                'subject': str(subject)[:60],
                'sender': str(sender)[:40],
                'level': level,
                'sub_category': str(sub_cat),
                'date': str(date),
                'owner': str(owner),
                'source_email': str(source),
            })
        wb.close()
        return emails[-max_count:]  # 最近 N 封
    except Exception as e:
        print(f'  [!] 读取报表失败: {e}')
        return []


def load_errors():
    """加载所有错误报告"""
    if not ERROR_FILE.exists():
        return []
    errors = []
    with open(ERROR_FILE, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    errors.append(json.loads(line))
                except json.JSONDecodeError:
                    continue
    return errors


def save_error(error: dict):
    """追加一条错误报告"""
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    with open(ERROR_FILE, 'a', encoding='utf-8') as f:
        f.write(json.dumps(error, ensure_ascii=False) + '\n')


def list_errors():
    """列出未处理的错误报告"""
    errors = load_errors()
    unprocessed = [e for e in errors if not e.get('processed', False)]
    if not unprocessed:
        print('\n  没有未处理的误分类报告。\n')
        return

    print(f'\n  未处理的误分类报告 ({len(unprocessed)} 条):')
    print('  ' + '-' * 60)
    for i, e in enumerate(unprocessed, 1):
        print(f'  {i:2d}. [{e.get("timestamp", "?")[:10]}]')
        print(f'      主题: {e.get("email_subject", "?")[:50]}')
        print(f'      错误: Level {e.get("wrong_level", "?")} ({e.get("wrong_sub_category", "?")})')
        print(f'      正确: Level {e.get("correct_level", "?")} ({e.get("correct_sub_category", "?")})')
        if e.get('reporter_note'):
            print(f'      备注: {e["reporter_note"]}')
        print()


def show_stats():
    """显示统计信息"""
    errors = load_errors()
    total = len(errors)
    unprocessed = len([e for e in errors if not e.get('processed', False)])
    processed = total - unprocessed

    print(f'\n  === 误分类报告统计 ===')
    print(f'  总报告数:   {total}')
    print(f'  未处理:     {unprocessed}')
    print(f'  已处理:     {processed}')

    # 提示词版本信息
    rules_path = Path('prompts/learned_rules.txt')
    if rules_path.exists():
        rules_text = rules_path.read_text(encoding='utf-8')
        version = 'v0'
        rule_count = 0
        for line in rules_text.splitlines():
            if line.startswith('# 版本: v'):
                version = line.split('v')[-1].strip()
            if line.strip() and not line.startswith('#') and line.strip()[0].isdigit():
                rule_count += 1
        print(f'  提示词规则: v{version} ({rule_count} 条)')
    else:
        print(f'  提示词规则: 未初始化')
    print()


def interactive_report():
    """交互式报错"""
    print('\n  === 分级报错 ===')
    print('  选择一封被错误分级的邮件，报告正确的级别。\n')

    emails = load_recent_emails()
    if not emails:
        print('  [!] 没有找到最近的邮件记录。请先运行一次单次检查。\n')
        return False

    # 显示邮件列表
    print(f'  最近 {len(emails)} 封邮件:')
    print('  ' + '-' * 70)
    for i, e in enumerate(emails, 1):
        level_str = f'L{e["level"]}'
        subj = e['subject'][:45]
        owner = e['owner'][:12]
        print(f'  {i:2d}. [{level_str}] {e["sub_category"][:6]:6s} | {owner:12s} | {subj}')
    print('  ' + '-' * 70)
    print(f'  {"0":>4s}. 手动输入邮件信息')
    print()

    # 选择邮件
    try:
        choice = input('  请选择邮件编号 (0 或序号): ').strip()
        if not choice:
            return False
        idx = int(choice)
        if idx == 0:
            return manual_report()
        if idx < 1 or idx > len(emails):
            print('  [!] 编号超出范围')
            return False
    except (ValueError, EOFError):
        return False

    selected = emails[idx - 1]
    print(f'\n  已选择: {selected["subject"][:50]}')
    print(f'  当前分级: Level {selected["level"]} ({selected["sub_category"]})')
    print(f'  负责人: {selected["owner"]}')
    print()

    # 输入正确级别
    print('  正确级别:')
    for lv, name in LEVEL_NAMES.items():
        print(f'    {lv}. {name}')
    print()

    try:
        correct_level = input('  请输入正确级别 (0-6): ').strip()
        correct_level = int(correct_level)
        if correct_level < 0 or correct_level > 6:
            print('  [!] 级别必须是 0-6')
            return False
    except (ValueError, EOFError):
        print('  [!] 无效输入')
        return False

    # 选择子类别
    sub_cats = SUB_CATS.get(correct_level, ['其他'])
    print(f'\n  {correct_level}级的子类别:')
    for i, sc in enumerate(sub_cats, 1):
        print(f'    {i}. {sc}')
    print()

    try:
        sc_choice = input('  请选择子类别编号: ').strip()
        sc_idx = int(sc_choice) - 1
        if sc_idx < 0 or sc_idx >= len(sub_cats):
            correct_sub_cat = sub_cats[0]
        else:
            correct_sub_cat = sub_cats[sc_idx]
    except (ValueError, EOFError):
        correct_sub_cat = sub_cats[0]

    # 备注
    try:
        note = input('  备注 (可选，按回车跳过): ').strip()
    except EOFError:
        note = ''

    # 构建错误记录
    now = datetime.now()
    error = {
        'id': f'mc_{now.strftime("%Y%m%d_%H%M%S")}',
        'timestamp': now.isoformat(timespec='seconds'),
        'email_subject': selected['subject'],
        'email_sender': selected['sender'],
        'email_body_preview': '',
        'email_date': selected['date'],
        'email_owner': selected['owner'],
        'email_source': selected['source_email'],
        'wrong_level': selected['level'],
        'wrong_sub_category': selected['sub_category'],
        'correct_level': correct_level,
        'correct_sub_category': correct_sub_cat,
        'reporter_note': note,
        'processed': False,
    }

    save_error(error)
    print(f'\n  ✅ 已记录误分类报告:')
    print(f'     Level {selected["level"]} ({selected["sub_category"]}) → Level {correct_level} ({correct_sub_cat})')
    print(f'     报告保存在: {ERROR_FILE}')
    print()
    return True


def manual_report():
    """手动输入邮件信息报错"""
    print('\n  --- 手动输入 ---')
    try:
        subject = input('  邮件主题: ').strip()
        sender = input('  发件人: ').strip()
        if not subject:
            print('  [!] 主题不能为空')
            return False

        print('\n  当前错误分级:')
        for lv, name in LEVEL_NAMES.items():
            print(f'    {lv}. {name}')
        wrong_level = int(input('  错误级别 (0-6): ').strip())

        print(f'\n  正确分级:')
        for lv, name in LEVEL_NAMES.items():
            print(f'    {lv}. {name}')
        correct_level = int(input('  正确级别 (0-6): ').strip())

        note = input('  备注 (可选): ').strip()

    except (ValueError, EOFError):
        print('  [!] 输入无效')
        return False

    now = datetime.now()
    error = {
        'id': f'mc_{now.strftime("%Y%m%d_%H%M%S")}',
        'timestamp': now.isoformat(timespec='seconds'),
        'email_subject': subject,
        'email_sender': sender,
        'email_body_preview': '',
        'email_date': now.strftime('%Y-%m-%d'),
        'email_owner': '',
        'email_source': '',
        'wrong_level': wrong_level,
        'wrong_sub_category': '',
        'correct_level': correct_level,
        'correct_sub_category': '',
        'reporter_note': note,
        'processed': False,
    }

    save_error(error)
    print(f'\n  ✅ 已记录误分类报告')
    return True


def main():
    os.chdir(os.path.dirname(os.path.abspath(__file__)))

    args = sys.argv[1:]

    if '--list' in args:
        list_errors()
    elif '--stats' in args:
        show_stats()
    elif '--help' in args or '-h' in args:
        print('''
  用法:
    python _report_error.py           交互式报错
    python _report_error.py --list    查看未处理的错误报告
    python _report_error.py --stats   查看统计信息
''')
    else:
        interactive_report()


if __name__ == '__main__':
    main()
